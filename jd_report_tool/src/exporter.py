from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .logger import logger
from .utils import PROJECT_ROOT, safe_text, unique_path

SPU_COLUMNS = ["时间", "spu", "货号", "曝光次数", "点击次数", "曝光点击率", "成交商品件数", "曝光转化率", "成交金额"]
SKU_COLUMNS = ["时间", "spu", "货号", "sku", "曝光次数", "点击次数", "曝光点击率", "成交商品件数", "曝光转化率", "成交金额", "类目排名"]
OUTPUT_COLUMNS = [
    "时间", "spu", "sku", "货号",
    "spu总访客", "spu总转化率", "spu订单数", "spu销售额",
    "sku排名", "曝光次数", "点击率", "订单数", "销售额",
]


def _format_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True)
    text_headers = {"spu", "sku", "货号"}
    percent_headers = {"曝光点击率", "曝光转化率", "spu总转化率", "点击率"}
    money_headers = {"成交金额", "到手价", "spu销售额", "销售额"}
    header_map = {cell.value: idx for idx, cell in enumerate(ws[1], start=1)}
    for header, idx in header_map.items():
        letter = get_column_letter(idx)
        max_len = len(str(header or ""))
        for row in range(2, ws.max_row + 1):
            cell = ws[f"{letter}{row}"]
            if header in text_headers:
                cell.number_format = "@"
                if cell.value is not None:
                    cell.value = safe_text(cell.value)
            elif header in percent_headers:
                cell.number_format = "0.00%"
            elif header in money_headers:
                cell.number_format = "#,##0.00"
            elif header == "时间":
                cell.number_format = "yyyy-mm-dd"
            max_len = max(max_len, len(safe_text(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 35)


def build_output_table(spu_out: pd.DataFrame, sku_out: pd.DataFrame) -> pd.DataFrame:
    spu_metrics = (
        spu_out.reindex(columns=["时间", "spu", "曝光次数", "曝光转化率", "成交商品件数", "成交金额"])
        .rename(columns={
            "曝光次数": "spu总访客",
            "曝光转化率": "spu总转化率",
            "成交商品件数": "spu订单数",
            "成交金额": "spu销售额",
        })
        .drop_duplicates(["时间", "spu"], keep="last")
    )
    sku_metrics = sku_out.reindex(columns=[
        "时间", "spu", "sku", "货号", "类目排名", "曝光次数", "曝光点击率", "成交商品件数", "成交金额",
    ]).rename(columns={
        "类目排名": "sku排名",
        "曝光点击率": "点击率",
        "成交商品件数": "订单数",
        "成交金额": "销售额",
    })
    output = sku_metrics.merge(spu_metrics, on=["时间", "spu"], how="left")
    return output.reindex(columns=OUTPUT_COLUMNS)


def export_summary(spu_df: pd.DataFrame, sku_df: pd.DataFrame, store_name: str, date_start: str, date_end: str,
                   output_dir: str | Path | None = None) -> Path:
    output = Path(output_dir) if output_dir else PROJECT_ROOT / "data" / "output"
    output.mkdir(parents=True, exist_ok=True)
    path = unique_path(output / f"{store_name}_SPU与SKU数据汇总_{date_start}_{date_end}.xlsx")
    spu_out = spu_df.reindex(columns=SPU_COLUMNS)
    sku_out = sku_df.reindex(columns=SKU_COLUMNS)
    output_out = build_output_table(spu_out, sku_out)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        spu_out.to_excel(writer, index=False, sheet_name="SPU数据汇总")
        sku_out.to_excel(writer, index=False, sheet_name="SKU数据汇总")
        output_out.to_excel(writer, index=False, sheet_name="输出表")
    wb = load_workbook(path)
    for name in ["SPU数据汇总", "SKU数据汇总", "输出表"]:
        _format_sheet(wb[name])
    wb.save(path)
    logger.info("汇总输出路径: %s", path)
    return path
